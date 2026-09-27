from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "EDURECOVER_CLASSROOM_BOARD_TEACHER_PRODUCT_FULL_DESIGN_20260928.md"
OUTPUT = ROOT / "artifacts" / "EduRecover_Classroom_Board_교사용_상품_전체설계_20260928.docx"

NAVY = "17365D"
BLUE = "2F5597"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F3F5F7"
BORDER = "D9D9D9"
TEXT = RGBColor(25, 35, 48)
MUTED = RGBColor(91, 103, 116)


def set_font(run, name: str = "Noto Sans KR", size: float | None = None, bold: bool | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = TEXT


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=130, bottom=80, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), BORDER)


def add_inline(paragraph, text: str, *, size: float = 10.8, bold=False, color=None):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, "Liberation Mono", size - 0.3)
            run.font.color.rgb = RGBColor(37, 75, 120)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=size, bold=True)
            if color:
                run.font.color.rgb = color
        else:
            run = paragraph.add_run(part)
            set_font(run, size=size, bold=bold)
            if color:
                run.font.color.rgb = color


def style_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Noto Sans KR"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans KR")
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.28

    title = doc.styles["Title"]
    title.font.name = "Noto Sans KR"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans KR")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(16)

    for name, size, before, after in (
        ("Heading 1", 16, 15, 7),
        ("Heading 2", 13, 11, 5),
        ("Heading 3", 11.5, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Noto Sans KR"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans KR")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("EduRecover Classroom Board")
    set_font(run, size=8.5, bold=True)
    run.font.color.rgb = MUTED

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("EduRecover 교사용 상품 전체설계  2026 09 28")
    set_font(run, size=8)
    run.font.color.rgb = MUTED


def add_markdown_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for r_idx, row in enumerate(rows):
        for c_idx in range(columns):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                shade_cell(cell, NAVY)
            elif r_idx % 2 == 0:
                shade_cell(cell, PALE_BLUE)
            else:
                shade_cell(cell, "FFFFFF")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            if c_idx == 0 and len(row) > 2:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            value = row[c_idx] if c_idx < len(row) else ""
            add_inline(paragraph, value, size=9.0, bold=r_idx == 0)
            for run in paragraph.runs:
                if r_idx == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        values = [item.strip() for item in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", value or "") for value in values):
            rows.append(values)
        i += 1
    return rows, i


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_document(doc)
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            add_markdown_table(doc, rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                paragraph = doc.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_inline(paragraph, text, size=24, bold=True)
            else:
                paragraph = doc.add_paragraph(style=f"Heading {level - 1}")
                add_inline(paragraph, text, size=16 if level == 2 else 13, bold=True)
            i += 1
            continue
        if re.match(r"^[-*]\s+", line):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.24)
            paragraph.paragraph_format.first_line_indent = Inches(-0.16)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.18
            add_inline(paragraph, re.sub(r"^[-*]\s+", "", line))
            i += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if numbered:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.26)
            paragraph.paragraph_format.first_line_indent = Inches(-0.26)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            add_inline(paragraph, f"{numbered.group(1)}. {numbered.group(2)}")
            i += 1
            continue
        paragraph_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt.startswith("#") or nxt.startswith("|") or re.match(r"^[-*]\s+", nxt) or re.match(r"^\d+\.\s+", nxt):
                break
            paragraph_lines.append(nxt)
            i += 1
        paragraph = doc.add_paragraph()
        if paragraph_lines[0].startswith("- 기준일:"):
            paragraph.paragraph_format.space_after = Pt(2)
        add_inline(paragraph, " ".join(paragraph_lines))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
