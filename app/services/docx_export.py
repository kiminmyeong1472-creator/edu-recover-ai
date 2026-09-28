"""Generate a clearly marked teacher-review draft and answer sheet."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


DOCX_FONT = "맑은 고딕"
CHOICE_LABELS = ("①", "②", "③", "④", "⑤", "⑥", "⑦", "⑧")


def _as_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _set_run_font(run: Any, *, size: float = 10.5, bold: bool | None = None) -> None:
    run.font.name = DOCX_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), DOCX_FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _add_text_paragraph(document: Document, text: str, *, bold: bool = False, size: float = 10.5):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    _set_run_font(run, size=size, bold=bold)
    return paragraph


def _choice_label(index: int) -> str:
    if index < len(CHOICE_LABELS):
        return CHOICE_LABELS[index]
    return f"{index + 1}."


def build_exam_docx(exam: dict[str, Any], items: list[dict[str, Any]]) -> bytes:
    """Export an editable draft, never an unverified student release.

    Client-side PASS/approval flags are not accepted as release evidence here.
    """
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    draft_notice = "교사용 검토 초안 · 학생 배포 전 문항·정답·근거 검증 필요"
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(header.add_run(draft_notice), size=9, bold=True)

    normal = document.styles["Normal"]
    normal.font.name = DOCX_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT)
    normal.font.size = Pt(10.5)

    title = _as_text(exam.get("name")) or "EduRecover 시험지"
    school_level = _as_text(exam.get("schoolLevel"))
    grade = _as_text(exam.get("grade"))
    exam_type = _as_text(exam.get("type"))
    scope = _as_text(exam.get("scope"))

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    _set_run_font(run, size=16, bold=True)
    _add_text_paragraph(document, draft_notice, bold=True, size=10)

    meta_parts = [part for part in (school_level, grade, exam_type) if part]
    if meta_parts:
        meta = document.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = meta.add_run(" · ".join(meta_parts))
        _set_run_font(run, size=10)

    if scope:
        _add_text_paragraph(document, f"출제 범위: {scope}", size=9.5)

    document.add_paragraph()

    for index, item in enumerate(items, start=1):
        question = _as_text(item.get("question"))
        passage = _as_text(item.get("passage"))
        choices = item.get("choices") if isinstance(item.get("choices"), list) else []

        question_p = document.add_paragraph()
        question_p.paragraph_format.space_before = Pt(8)
        question_p.paragraph_format.space_after = Pt(4)
        run = question_p.add_run(f"{index}. {question or '문항'}")
        _set_run_font(run, size=11, bold=True)

        if passage:
            passage_p = _add_text_paragraph(document, passage, size=10.5)
            passage_p.paragraph_format.space_after = Pt(4)

        for choice_index, choice in enumerate(choices):
            text = _as_text(choice)
            if not text:
                continue
            choice_p = _add_text_paragraph(
                document,
                f"{_choice_label(choice_index)} {text}",
                size=10.5,
            )
            choice_p.paragraph_format.left_indent = Inches(0.12)
            choice_p.paragraph_format.space_after = Pt(1.5)

    document.add_page_break()

    answer_heading = document.add_paragraph()
    answer_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = answer_heading.add_run("정답 · 해설")
    _set_run_font(run, size=16, bold=True)

    for index, item in enumerate(items, start=1):
        answer = _as_text(item.get("answer"))
        explanation = _as_text(item.get("explanation"))
        evidence = _as_text(item.get("evidence"))

        answer_p = document.add_paragraph()
        answer_p.paragraph_format.space_before = Pt(8)
        answer_p.paragraph_format.space_after = Pt(3)
        run = answer_p.add_run(f"{index}. 정답: {answer or '-'}")
        _set_run_font(run, size=11, bold=True)

        if explanation:
            _add_text_paragraph(document, f"해설: {explanation}", size=10.5)
        if evidence:
            evidence_p = _add_text_paragraph(document, f"근거: {evidence}", size=9.5)
            evidence_p.paragraph_format.left_indent = Inches(0.12)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
