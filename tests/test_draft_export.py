from io import BytesIO

from docx import Document

from app.services.docx_export import build_exam_docx


def test_client_pass_flags_do_not_turn_unverified_export_into_final_exam():
    raw = build_exam_docx({"name": "검토본", "teacherApproved": True,
                           "crossVerificationState": "PASS"},
                          [{"question": "문항", "choices": ["가", "나"],
                            "answer": "가", "release_ready": True}])
    doc = Document(BytesIO(raw))
    assert "교사용 검토 초안" in " ".join(p.text for p in doc.paragraphs)
    assert "학생 배포 전" in doc.sections[0].header.paragraphs[0].text
